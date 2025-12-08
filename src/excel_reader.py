"""Excel reader for extracting invoice data using openpyxl.

Reads invoice data from the 'account credit vendor' sheet in company_data.xlsx
and returns unique Invoice objects compatible with QuickBooks sync.
"""

from __future__ import annotations
from pathlib import Path
from datetime import datetime, date
from typing import List
from openpyxl import load_workbook
from src.models import Invoice


def extract_invoices(
    workbook_path: Path, sheet_name: str = "account credit vendor"
) -> List[Invoice]:
    """Extract unique, well-formed invoice records from Excel.

    Args:
        workbook_path: Path to the Excel file.
        sheet_name: Worksheet name to read from.

    Returns:
        List of unique Invoice objects.
    """
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(filename=workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{sheet_name}' not found. Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)

        headers = [str(h).strip() if h else "" for h in next(rows, [])]
        header_index = {name: i for i, name in enumerate(headers)}

        invoices: List[Invoice] = []
        seen_keys = set()  # Deduplicate by (customer, invoice_number)

        for row in rows:
            if not any(row):
                continue  # skip empty rows

            def _value(column_name: str):
                """Safe header lookup with fallback."""
                idx = header_index.get(column_name)
                if idx is None or idx >= len(row):
                    return None
                return row[idx]

            record_id = str(_value("Child ID") or "").strip()
            customer = str(_value("Customer") or "").strip()
            invoice_number = str(
                _value("Invoice Number") or _value("Invoice Num") or ""
            ).strip()
            invoice_date = _value("Invoice Date")

            # Convert datetime to date safely
            if isinstance(invoice_date, datetime):
                invoice_date = invoice_date.date()
            elif not isinstance(invoice_date, date) or invoice_date is None:
                invoice_date = date.today()  # fallback for QuickBooks

            try:
                invoice_amount = float(_value("Invoice Amount") or 0)
            except (TypeError, ValueError):
                invoice_amount = 0.0

            if not invoice_number or not customer:
                continue  # skip incomplete records

            key = (customer.lower(), invoice_number)
            if key in seen_keys:
                continue  # skip duplicates
            seen_keys.add(key)

            invoices.append(
                Invoice(
                    record_id=record_id,
                    customer=customer,
                    invoice_number=invoice_number,
                    invoice_date=invoice_date,
                    invoice_amount=invoice_amount,
                    source="excel",
                )
            )

    finally:
        workbook.close()

    print(f" Extracted {len(invoices)} unique invoices from Excel")
    return invoices


__all__ = ["extract_invoices"]


if __name__ == "__main__":
    """Manual test for Excel extraction."""
    import sys

    try:
        excel_path = (
            Path(sys.argv[1]) if len(sys.argv) > 1 else Path("company_data.xlsx")
        )
        invoices = extract_invoices(excel_path)
        for inv in invoices:
            print(inv)
    except Exception as e:
        print(f"❌ Error: {e}")
        print("Usage: python -m src.excel_reader <path-to-workbook.xlsx>")
        sys.exit(1)
